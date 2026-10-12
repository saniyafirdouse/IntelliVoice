"""
Admin Panel + Analytics tests. Uses a FAKE in-memory MongoDB, so your real
Atlas data is never touched and no internet is needed.

One-time setup:   pip install mongomock-motor
Run from backend/:  python -m tests.test_admin
"""

import io
import json
import os
from datetime import datetime, timedelta, timezone

os.environ.setdefault("SECRET_KEY", "test-secret-key-for-local-tests-only-0123456789")

import asyncio

from fastapi import FastAPI
from fastapi.testclient import TestClient
from mongomock_motor import AsyncMongoMockClient
from openpyxl import load_workbook

import database.connection as conn
from engines import analytics_engine
from engines.knowledge_engine import retrieve_knowledge
from routes.admin import router as admin_router
from routes.auth import router as auth_router
from utils.auth import hash_password
from utils.data_import import today_ist

SEED = os.path.join(os.path.dirname(__file__), "..", "database", "seed_data", "fees.json")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def run(coro):
    return asyncio.get_event_loop().run_until_complete(coro)


def fresh_app():
    """New fake database with the real fees seed data + one admin + one student."""
    conn.db = AsyncMongoMockClient()["intellivoice_test"]
    with open(SEED, encoding="utf-8") as f:
        run(conn.db["fees"].insert_many(json.load(f)))
    run(conn.db["users"].insert_many([
        {"name": "Admin", "email": "admin@svit.test", "role": "admin", "password_hash": hash_password("admin123!")},
        {"name": "Stu", "email": "stu@svit.test", "role": "student", "password_hash": hash_password("stud123!")},
    ]))
    app = FastAPI()
    app.include_router(auth_router, prefix="/auth")
    app.include_router(admin_router, prefix="/api/admin")
    return TestClient(app)


def admin_headers(client):
    r = client.post("/auth/login", json={"email": "admin@svit.test", "password": "admin123!", "role": "admin"})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['token']}"}


def edited_template(client, h, edits=None, extra_rows=None):
    """Download the template, change some cells, return xlsx bytes."""
    r = client.get("/api/admin/datasets/fees/template", headers=h)
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["fees"]
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        vals = dict(zip(header, [c.value for c in row]))
        for match, changes in (edits or []):
            if all(vals.get(k) == v for k, v in match.items()):
                for k, v in changes.items():
                    row[header.index(k)].value = v
    for extra in (extra_rows or []):
        ws.append([extra.get(f) for f in header])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(client, h, content, filename="fees.xlsx", mime=XLSX):
    return client.post("/api/admin/datasets/fees/upload", headers=h,
                       files={"file": (filename, content, mime)})


CSE_MGMT = {"branch": "CSE", "quota": "Management"}


def test_login_and_protection():
    c = fresh_app()
    assert c.get("/api/admin/datasets").status_code == 401
    bad = c.post("/auth/login", json={"email": "admin@svit.test", "password": "wrong", "role": "admin"})
    assert bad.status_code == 401
    stu = c.post("/auth/login", json={"email": "stu@svit.test", "password": "stud123!", "role": "student"}).json()
    r = c.get("/api/admin/datasets", headers={"Authorization": f"Bearer {stu['token']}"})
    assert r.status_code == 403
    su = c.post("/auth/signup", json={"name": "P", "email": "p@x.com", "password": "secret1"})
    assert su.status_code == 201 and su.json()["user"]["role"] == "student"
    stored = run(conn.db["users"].find_one({"email": "p@x.com"}))
    assert stored["password_hash"] != "secret1"
    print("PASS 1: login works, passwords hashed, admin pages blocked for students/no login")


def test_template_prefilled():
    c = fresh_app()
    h = admin_headers(c)
    ds = c.get("/api/admin/datasets", headers=h).json()["datasets"][0]
    assert ds["name"] == "fees" and ds["records"] == 24
    wb = load_workbook(io.BytesIO(c.get("/api/admin/datasets/fees/template", headers=h).content))
    assert wb["fees"].max_row == 25 and "Instructions" in wb.sheetnames
    print("PASS 2: template downloads with all 24 current fee records + instructions sheet")


def test_validation_errors_block_confirm():
    c = fresh_app()
    h = admin_headers(c)
    content = edited_template(c, h, extra_rows=[
        {"branch": "CSEE", "exam": "KCET", "academic_year": "2026-27", "annual_fee": 100000,
         "semester_fee": 100000, "payment_mode": "One-time", "installments": 1},
        {"branch": "ECE", "exam": "KCET", "quota": "NRI", "academic_year": "2026-27", "annual_fee": "abc",
         "semester_fee": 100000, "payment_mode": "One-time", "installments": 1},
    ])
    p = upload(c, h, content).json()
    msgs = " | ".join(e["message"] for e in p["errors"])
    assert "branch' = 'CSEE' is not valid" in msgs, msgs
    assert "annual_fee' must be a number" in msgs, msgs
    assert p["can_confirm"] is False
    r = c.post(f"/api/admin/uploads/{p['upload_id']}/confirm", headers=h)
    assert r.status_code == 400
    print("PASS 3: bad rows rejected with clear row-numbered messages; confirm blocked")


def test_preview_confirm_undo():
    c = fresh_app()
    h = admin_headers(c)
    content = edited_template(
        c, h,
        edits=[(CSE_MGMT, {"annual_fee": 350000, "semester_fee": 175000})],
        extra_rows=[{"branch": "ECE", "quota": "NRI", "academic_year": "2026-27", "annual_fee": 500000,
                     "semester_fee": 250000, "payment_mode": "Semester-wise", "installments": 2}],
    )
    p = upload(c, h, content).json()
    assert p["counts"]["updated"] == 1 and p["counts"]["new"] == 1 and p["counts"]["unchanged"] == 23, p["counts"]
    upd = next(ch for ch in p["changes"] if ch["action"] == "update")
    assert {"field": "annual_fee", "old": 325000, "new": 350000} in upd["fields"]
    # Nothing saved yet
    doc = run(conn.db["fees"].find_one(CSE_MGMT))
    assert doc["annual_fee"] == 325000

    r = c.post(f"/api/admin/uploads/{p['upload_id']}/confirm", headers=h)
    assert r.status_code == 200 and r.json()["status"] == "applied", r.text
    doc = run(conn.db["fees"].find_one(CSE_MGMT))
    assert doc["annual_fee"] == 350000 and doc["last_updated"] == today_ist()
    k = run(retrieve_knowledge("fees", {"branch": "CSE", "quota": "Management"}))
    assert k["data"]["annual_fee"] == 350000  # Engine 3 sees the new value
    assert run(conn.db["fees"].count_documents({})) == 25
    assert c.post(f"/api/admin/uploads/{p['upload_id']}/confirm", headers=h).status_code == 409

    hist = c.get("/api/admin/uploads", headers=h).json()["uploads"]
    assert hist[0]["status"] == "applied"

    r = c.post(f"/api/admin/uploads/{p['upload_id']}/undo", headers=h)
    assert r.status_code == 200, r.text
    doc = run(conn.db["fees"].find_one(CSE_MGMT))
    assert doc["annual_fee"] == 325000 and doc["last_updated"] == "2026-04-10"
    assert run(conn.db["fees"].count_documents({})) == 24
    print("PASS 4: preview shows exact changes, confirm saves with today's date, undo restores everything")


def test_warnings_and_aliases_csv():
    c = fresh_app()
    h = admin_headers(c)
    csv_text = (
        "Branch,Exam,Quota,Academic Year,Annual Fee,Semester Fee,Early Bird Fee,Payment Mode,Installments\n"
        "computer science,,management quota,2026-27,\"3,50,000\",175000,275000,semester wise,2\n"
        "CSE,KCET,,2026-27,12000,12000,,One-time,1\n"
    )
    p = upload(c, h, csv_text.encode("utf-8"), "fees.csv", "text/csv").json()
    assert not p["errors"], p["errors"]
    assert p["counts"]["updated"] == 2 and p["counts"]["not_in_file"] == 22, p["counts"]
    warns = " | ".join(w for ch in p["changes"] for w in ch["warnings"])
    assert "annual_fee changes by 90%" in warns, warns  # the missing-zero catch
    print("PASS 5: CSV + aliases accepted ('computer science' → CSE), missing-zero mistake flagged")


def test_undo_conflict_and_stale_preview():
    c = fresh_app()
    h = admin_headers(c)
    p1 = upload(c, h, edited_template(c, h, edits=[(CSE_MGMT, {"annual_fee": 330000, "semester_fee": 165000})])).json()
    # a second preview made BEFORE p1 is confirmed, touching the same record
    p_stale = upload(c, h, edited_template(c, h, edits=[(CSE_MGMT, {"annual_fee": 331000})])).json()
    assert c.post(f"/api/admin/uploads/{p1['upload_id']}/confirm", headers=h).status_code == 200
    r = c.post(f"/api/admin/uploads/{p_stale['upload_id']}/confirm", headers=h)
    assert r.status_code == 409 and r.json()["records"], r.text

    p2 = upload(c, h, edited_template(c, h, edits=[(CSE_MGMT, {"annual_fee": 340000, "semester_fee": 170000})])).json()
    assert c.post(f"/api/admin/uploads/{p2['upload_id']}/confirm", headers=h).status_code == 200
    assert c.post(f"/api/admin/uploads/{p1['upload_id']}/undo", headers=h).status_code == 409
    assert c.post(f"/api/admin/uploads/{p2['upload_id']}/undo", headers=h).status_code == 200
    assert c.post(f"/api/admin/uploads/{p1['upload_id']}/undo", headers=h).status_code == 200
    assert run(conn.db["fees"].find_one(CSE_MGMT))["annual_fee"] == 325000
    print("PASS 6: outdated previews and out-of-order undos are blocked, nothing gets overwritten")


def test_analytics():
    c = fresh_app()
    h = admin_headers(c)
    db = conn.db

    def intel(intent, zone="confident", category="Admissions", collection="fees", entities=None):
        return {"intent": intent, "category": category, "confidence": 0.9, "confidence_zone": zone,
                "emotion": "neutral", "entities": entities or {}, "response_mode": "informational",
                "knowledge_request": {"collection": collection}, "needs_escalation": False}

    found = {"found": True}
    nodata = {"found": False, "reason": "no_matching_document"}
    calls = [
        ("CSE fee kitni hai?", "hi", intel("ask_fees_branch"), found),
        ("What is the CSE fee?", "en", intel("ask_fees_branch"), found),
        ("Bus route to Yelahanka?", "en", intel("ask_transport", collection="transport",
                                                  entities={"campus_facility": "transport"}), nodata),
        ("bus route to yelahanka?", "kn", intel("ask_transport", collection="transport",
                                                  entities={"campus_facility": "transport"}), nodata),
        ("Who won the cricket match?", "en", intel("fallback", zone="out_of_scope"), None),
        ("who won the  cricket match?", "en", intel("fallback", zone="out_of_scope"), None),
    ]
    for text, lang, it, kr in calls:
        run(analytics_engine.log_query(db, session_id=f"s-{lang}", text=text,
                                       language={"primary": lang, "is_code_switched": lang == "hi"},
                                       intelligence_result=it, knowledge_result=kr, response_time_ms=1200))
    # logging never crashes even with garbage input
    run(analytics_engine.log_query(None, session_id="x", text="x", language=None, intelligence_result={}))

    d = c.get("/api/admin/dashboard?days=7", headers=h).json()
    assert d["total_queries"] == 6 and d["status"]["answered"] == 2 and d["status"]["no_data"] == 2
    assert d["answered_rate"] == 33.3 and d["top_intents"][0] == {"intent": "ask_fees_branch", "count": 2}
    assert sum(x["count"] for x in d["busiest_hours"]) == 6 and len(d["daily_queries"]) == 7

    u = c.get("/api/admin/unanswered", headers=h).json()
    assert u["missing_data"][0]["intent"] == "ask_transport" and u["missing_data"][0]["count"] == 2
    assert u["out_of_scope"][0]["count"] == 2  # same question, different spacing/case
    ids = u["missing_data"][0]["log_ids"]
    assert c.post("/api/admin/unanswered/resolve", headers=h, json={"log_ids": ids}).json()["resolved"] == 2
    u2 = c.get("/api/admin/unanswered", headers=h).json()
    assert u2["missing_data"] == [] and u2["total_unanswered"] == 2
    print("PASS 7: questions logged; dashboard + 'couldn't answer' list correct; resolve works")


if __name__ == "__main__":
    asyncio.set_event_loop(asyncio.new_event_loop())
    test_login_and_protection()
    test_template_prefilled()
    test_validation_errors_block_confirm()
    test_preview_confirm_undo()
    test_warnings_and_aliases_csv()
    test_undo_conflict_and_stale_preview()
    test_analytics()
    print("\nAll 7 Admin Panel tests passed.")
