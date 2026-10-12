"""
utils/data_import.py — Excel/CSV upload → validate → preview → confirm → undo

The flow (each step is one admin API call):

  1. build_template()  Download an Excel file pre-filled with the CURRENT data,
                       so the admin edits real values instead of retyping.
  2. read_sheet() + validate_rows()
                       Read the uploaded file and check every row
                       ("Row 5: branch 'CSEE' is not valid").
  3. diff_against_db() Compare with MongoDB and list exactly what will change
                       ("annual_fee 1,20,000 → 1,35,000"). Saved as a PENDING
                       upload batch. Nothing in the real data changes yet.
  4. confirm_batch()   Admin clicks Confirm → changes are written, each changed
                       record gets last_updated = today.
  5. undo_batch()      Puts every record back exactly as it was before.

The upload batches (collection "admin_uploads") are the version history:
each one stores the old and new version of every record it touched.

Rows in MongoDB that are NOT in the uploaded file are left untouched,
so an upload can never delete data by accident.
"""

import csv
import io
import re
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo

from bson import ObjectId
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from utils.datasets import BIG_CHANGE_RATIO

UPLOADS = "admin_uploads"
IST = ZoneInfo("Asia/Kolkata")
PENDING_HOURS = 24          # a preview must be confirmed within a day
MAX_ROWS = 2000


def today_ist() -> str:
    return datetime.now(IST).strftime("%Y-%m-%d")


def _norm(s) -> str:
    """'Annual Fee ' / 'annual_fee' / 'ANNUAL-FEE' -> 'annualfee'."""
    return re.sub(r"[\s_\-]+", "", str(s or "")).lower()


# ═══════════════════════════════════════════════════════════════════════════
# 1. Template
# ═══════════════════════════════════════════════════════════════════════════
def build_template(dataset_name: str, dataset: dict, current_docs: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = dataset_name
    fields = [c["field"] for c in dataset["columns"]]

    ws.append(fields)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    ws.freeze_panes = "A2"

    key = lambda d: tuple(str(d.get(k) or "") for k in dataset["key_fields"])
    for doc in sorted(current_docs, key=key):
        ws.append([doc.get(f) for f in fields])

    for i, col in enumerate(dataset["columns"], start=1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = max(14, len(col["field"]) + 4)
        if col["type"] == "choice":  # dropdown in Excel
            options = ",".join(col["choices"].keys())
            if len(options) < 250:  # Excel's limit for inline lists
                dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
                dv.error = f"Choose one of: {options}"
                ws.add_data_validation(dv)
                dv.add(f"{letter}2:{letter}{MAX_ROWS + 1}")

    # Instructions sheet
    info = wb.create_sheet("Instructions")
    info.append([f"{dataset['title']}: how to update"])
    info["A1"].font = Font(bold=True, size=13)
    for line in [
        "Edit the values in the first sheet, or add new rows. Do not rename the column headers.",
        "Rows you delete from this file are NOT deleted from the system (it is always safe).",
        "After uploading you will see a preview of every change. Nothing is saved until you click Confirm.",
        "'last_updated' is filled automatically, so you don't need to add it.",
        "",
    ]:
        info.append([line])
    info.append(["Column", "Required?", "Allowed values / format"])
    for c in info[info.max_row]:
        c.font = Font(bold=True)
    for col in dataset["columns"]:
        allowed = ", ".join(col["choices"].keys()) if col["type"] == "choice" else col.get("help", "")
        info.append([col["field"], "Yes" if col.get("required") else "Optional", allowed])
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 12
    info.column_dimensions["C"].width = 90

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ═══════════════════════════════════════════════════════════════════════════
# 2. Read + validate
# ═══════════════════════════════════════════════════════════════════════════
class SheetError(ValueError):
    """The file itself is unusable (wrong type, missing columns, empty...)."""


def read_sheet(filename: str, content: bytes, dataset_name: str) -> tuple[list[str], list[tuple[int, list]]]:
    """Returns (header, [(excel_row_number, values), ...])."""
    name = (filename or "").lower()
    if name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise SheetError("CSV file must be saved as UTF-8. In Excel use 'Save As → CSV UTF-8'.")
        rows = list(csv.reader(io.StringIO(text)))
    elif name.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:
            raise SheetError("Could not open this Excel file. Please upload the .xlsx template.")
        ws = wb[dataset_name] if dataset_name in wb.sheetnames else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        raise SheetError("Please upload an .xlsx or .csv file.")

    if not rows or not any(rows[0]):
        raise SheetError("The file is empty. The first row must contain the column headers.")
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    data = [(i, r) for i, r in enumerate(rows[1:], start=2)
            if any(v not in (None, "") and str(v).strip() for v in r)]
    if not data:
        raise SheetError("No data rows found under the headers.")
    if len(data) > MAX_ROWS:
        raise SheetError(f"Too many rows ({len(data)}). Maximum is {MAX_ROWS} per upload.")
    return header, data


def _parse_int(value):
    if isinstance(value, bool):
        raise ValueError
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if value != int(value):
            raise ValueError
        return int(value)
    s = re.sub(r"[₹,\s]|/-$|rs\.?", "", str(value).strip(), flags=re.I)
    if re.fullmatch(r"\d+(\.0+)?", s):
        return int(float(s))
    raise ValueError


def _parse_choice(value, choices: dict):
    v = str(value).strip().lower()
    for canonical, aliases in choices.items():
        if v == canonical.lower() or v in [a.lower() for a in aliases]:
            return canonical
    raise ValueError


def _parse_cell(value, col: dict):
    """Returns (parsed_value_or_None, error_message_or_None)."""
    blank = value is None or str(value).strip() == ""
    if blank:
        return None, (f"'{col['field']}' is required" if col.get("required") else None)

    try:
        if col["type"] == "int":
            n = _parse_int(value)
            lo, hi = col.get("min"), col.get("max")
            if (lo is not None and n < lo) or (hi is not None and n > hi):
                return None, f"'{col['field']}' = {n:,} looks wrong (expected between {lo:,} and {hi:,})"
            return n, None
        if col["type"] == "choice":
            return _parse_choice(value, col["choices"]), None
        s = str(value).strip()
        if col.get("pattern") and not re.fullmatch(col["pattern"], s):
            return None, f"'{col['field']}' = '{s}' has the wrong format ({col.get('help', '')})"
        return s, None
    except ValueError:
        if col["type"] == "int":
            return None, f"'{col['field']}' must be a number, got '{value}'"
        allowed = ", ".join(col["choices"].keys())
        return None, f"'{col['field']}' = '{value}' is not valid. Allowed: {allowed}"


def validate_rows(header: list[str], data: list[tuple[int, list]], dataset: dict) -> dict:
    """
    Returns {"rows": [{"row": n, "values": {...}, "warnings": [...]}],
             "errors": [{"row": n, "message": ...}], "file_warnings": [...]}
    """
    columns = dataset["columns"]
    by_norm = {_norm(c["field"]): c for c in columns}
    col_index, file_warnings = {}, []
    for i, h in enumerate(header):
        c = by_norm.get(_norm(h))
        if c:
            col_index[c["field"]] = i
        elif h and _norm(h) not in ("lastupdated", "id"):
            file_warnings.append(f"Column '{h}' is not used and will be ignored")

    missing = [c["field"] for c in columns if c.get("required") and c["field"] not in col_index]
    if missing:
        raise SheetError(f"Missing column(s): {', '.join(missing)}. Please use the downloaded template.")

    rows, errors, seen_keys = [], [], {}
    for row_no, raw in data:
        values, row_errors = {}, []
        for col in columns:
            idx = col_index.get(col["field"])
            cell = raw[idx] if idx is not None and idx < len(raw) else None
            parsed, err = _parse_cell(cell, col)
            if err:
                row_errors.append(err)
            values[col["field"]] = parsed

        warnings = []
        if not row_errors and dataset.get("row_check"):
            extra_err, warnings = dataset["row_check"](values)
            row_errors += extra_err

        if not row_errors:
            key = tuple(values.get(k) for k in dataset["key_fields"])
            if key in seen_keys:
                row_errors.append(f"Same record as row {seen_keys[key]} (duplicate)")
            else:
                seen_keys[key] = row_no

        for e in row_errors:
            errors.append({"row": row_no, "message": f"Row {row_no}: {e}"})
        if not row_errors:
            rows.append({"row": row_no, "values": values, "warnings": warnings})

    return {"rows": rows, "errors": errors, "file_warnings": file_warnings}


# ═══════════════════════════════════════════════════════════════════════════
# 3. Diff against the database
# ═══════════════════════════════════════════════════════════════════════════
def record_key(doc: dict, key_fields: list[str]) -> tuple:
    return tuple(doc.get(k) or None for k in key_fields)


def key_label(doc: dict, key_fields: list[str]) -> str:
    """Human label like 'CSE · Management · 2026-27'."""
    return " · ".join(str(doc[k]) for k in key_fields if doc.get(k))


def diff_against_db(valid_rows: list[dict], current_docs: list[dict], dataset: dict) -> dict:
    key_fields = dataset["key_fields"]
    fields = [c["field"] for c in dataset["columns"]]
    int_fields = {c["field"] for c in dataset["columns"] if c["type"] == "int"}
    existing = {record_key(d, key_fields): d for d in current_docs}

    changes, unchanged = [], 0
    for r in valid_rows:
        new_vals = r["values"]
        old = existing.get(record_key(new_vals, key_fields))
        label = key_label(new_vals, key_fields)
        warnings = list(r["warnings"])

        if old is None:
            new_doc = {f: v for f, v in new_vals.items() if v is not None}
            changes.append({"action": "new", "row": r["row"], "label": label,
                            "record_id": ObjectId(), "old": None, "new": new_doc,
                            "fields": [{"field": f, "old": None, "new": v} for f, v in new_doc.items()],
                            "warnings": warnings})
            continue

        field_changes = []
        for f in fields:
            o, n = old.get(f), new_vals.get(f)
            if o == n or (o in (None, "") and n is None):
                continue
            field_changes.append({"field": f, "old": o, "new": n})
            if f in int_fields and isinstance(o, (int, float)) and o and n is not None:
                ratio = abs(n - o) / o
                if ratio > BIG_CHANGE_RATIO:
                    warnings.append(f"{f} changes by {ratio:.0%} ({o:,} → {n:,}). Is that correct?")
            if n is None:
                warnings.append(f"{f} will be removed (cell left empty)")

        if not field_changes:
            unchanged += 1
            continue

        new_doc = {k: v for k, v in old.items() if k not in fields}  # keep _id, last_updated, extra fields
        new_doc.update({f: v for f, v in new_vals.items() if v is not None})
        changes.append({"action": "update", "row": r["row"], "label": label,
                        "record_id": old["_id"], "old": old, "new": new_doc,
                        "fields": field_changes, "warnings": warnings})

    file_keys = {record_key(r["values"], key_fields) for r in valid_rows}
    not_in_file = sum(1 for k in existing if k not in file_keys)
    return {"changes": changes, "unchanged": unchanged, "not_in_file": not_in_file}


# ═══════════════════════════════════════════════════════════════════════════
# DB operations (create preview batch / confirm / cancel / undo)
# ═══════════════════════════════════════════════════════════════════════════
class BatchError(Exception):
    def __init__(self, status_code: int, message: str, details=None):
        super().__init__(message)
        self.status_code, self.message, self.details = status_code, message, details


def _clean_doc(doc):
    """Make a record JSON-safe for API responses (ObjectId → str)."""
    if isinstance(doc, dict):
        return {k: _clean_doc(v) for k, v in doc.items()}
    if isinstance(doc, list):
        return [_clean_doc(v) for v in doc]
    if isinstance(doc, ObjectId):
        return str(doc)
    if isinstance(doc, datetime):
        return doc.isoformat()
    return doc


def batch_summary(batch: dict, include_changes: bool = True) -> dict:
    changes = batch.get("changes", [])
    out = {
        "upload_id": str(batch["_id"]),
        "dataset": batch["dataset"],
        "filename": batch.get("filename"),
        "status": batch["status"],
        "uploaded_by": batch.get("uploaded_by"),
        "uploaded_at": batch.get("uploaded_at"),
        "applied_at": batch.get("applied_at"),
        "undone_at": batch.get("undone_at"),
        "counts": {
            "new": sum(c["action"] == "new" for c in changes),
            "updated": sum(c["action"] == "update" for c in changes),
            "unchanged": batch.get("unchanged", 0),
            "not_in_file": batch.get("not_in_file", 0),
            "errors": len(batch.get("errors", [])),
            "warnings": sum(len(c["warnings"]) for c in changes),
        },
        "can_confirm": batch["status"] == "pending" and not batch.get("errors") and bool(changes),
        "errors": batch.get("errors", []),
        "file_warnings": batch.get("file_warnings", []),
    }
    if include_changes:
        out["changes"] = [{k: c[k] for k in ("action", "row", "label", "fields", "warnings")} for c in changes]
    return _clean_doc(out)


async def create_preview(db, dataset_name, dataset, filename, content, admin_email) -> dict:
    header, data = read_sheet(filename, content, dataset_name)
    checked = validate_rows(header, data, dataset)
    current = await db[dataset["collection"]].find({}).to_list(length=None)
    diff = diff_against_db(checked["rows"], current, dataset)

    batch = {
        "dataset": dataset_name,
        "collection": dataset["collection"],
        "filename": filename,
        "status": "pending",
        "uploaded_by": admin_email,
        "uploaded_at": datetime.now(timezone.utc),
        "errors": checked["errors"],
        "file_warnings": checked["file_warnings"],
        "changes": diff["changes"],
        "unchanged": diff["unchanged"],
        "not_in_file": diff["not_in_file"],
    }
    result = await db[UPLOADS].insert_one(batch)
    batch["_id"] = result.inserted_id
    return batch_summary(batch)


async def _get_batch(db, upload_id: str) -> dict:
    try:
        oid = ObjectId(upload_id)
    except Exception:
        raise BatchError(404, "Upload not found.")
    batch = await db[UPLOADS].find_one({"_id": oid})
    if not batch:
        raise BatchError(404, "Upload not found.")
    return batch


def _same(a: dict | None, b: dict | None) -> bool:
    return _clean_doc(a) == _clean_doc(b)


async def confirm_batch(db, upload_id: str, admin_email: str) -> dict:
    batch = await _get_batch(db, upload_id)
    if batch["status"] != "pending":
        raise BatchError(409, f"This upload is already {batch['status']}.")
    if batch.get("errors"):
        raise BatchError(400, "Fix the errors in the file and upload it again before confirming.")
    if not batch["changes"]:
        raise BatchError(400, "There is nothing to change in this upload.")
    uploaded_at = batch["uploaded_at"]
    if uploaded_at.tzinfo is None:
        uploaded_at = uploaded_at.replace(tzinfo=timezone.utc)
    if datetime.now(timezone.utc) - uploaded_at > timedelta(hours=PENDING_HOURS):
        await db[UPLOADS].update_one({"_id": batch["_id"]}, {"$set": {"status": "expired"}})
        raise BatchError(409, "This preview is older than 24 hours. Please upload the file again.")

    coll = db[batch["collection"]]

    # Safety: has someone changed these records since the preview was made?
    conflicts = []
    for c in batch["changes"]:
        current = await coll.find_one({"_id": c["record_id"]})
        if c["action"] == "update" and not _same(current, c["old"]):
            conflicts.append(c["label"])
        if c["action"] == "new" and current is not None:
            conflicts.append(c["label"])
    if conflicts:
        raise BatchError(409, "Some records changed after this preview was made. Please upload the file again.",
                         conflicts)

    stamp = today_ist()
    for c in batch["changes"]:
        new_doc = dict(c["new"])
        new_doc["_id"] = c["record_id"]
        new_doc["last_updated"] = stamp
        c["new"] = new_doc
        if c["action"] == "new":
            await coll.insert_one(new_doc)
        else:
            await coll.replace_one({"_id": c["record_id"]}, new_doc)

    batch["status"], batch["applied_at"], batch["applied_by"] = "applied", datetime.now(timezone.utc), admin_email
    await db[UPLOADS].update_one({"_id": batch["_id"]}, {"$set": {
        "status": "applied", "applied_at": batch["applied_at"], "applied_by": admin_email,
        "changes": batch["changes"]}})
    return batch_summary(batch)


async def cancel_batch(db, upload_id: str) -> dict:
    batch = await _get_batch(db, upload_id)
    if batch["status"] != "pending":
        raise BatchError(409, f"Only a pending upload can be cancelled (this one is {batch['status']}).")
    await db[UPLOADS].update_one({"_id": batch["_id"]}, {"$set": {"status": "cancelled"}})
    batch["status"] = "cancelled"
    return batch_summary(batch, include_changes=False)


async def undo_batch(db, upload_id: str, admin_email: str) -> dict:
    batch = await _get_batch(db, upload_id)
    if batch["status"] != "applied":
        raise BatchError(409, f"Only an applied upload can be undone (this one is {batch['status']}).")

    coll = db[batch["collection"]]
    conflicts = []
    for c in batch["changes"]:
        current = await coll.find_one({"_id": c["record_id"]})
        if not _same(current, c["new"]):
            conflicts.append(c["label"])
    if conflicts:
        raise BatchError(409, "Can't undo: these records were changed again by a later upload. "
                              "Undo the later upload first.", conflicts)

    for c in batch["changes"]:
        if c["action"] == "new":
            await coll.delete_one({"_id": c["record_id"]})
        else:
            await coll.replace_one({"_id": c["record_id"]}, c["old"])

    now = datetime.now(timezone.utc)
    await db[UPLOADS].update_one({"_id": batch["_id"]}, {"$set": {
        "status": "undone", "undone_at": now, "undone_by": admin_email}})
    batch["status"], batch["undone_at"] = "undone", now
    return batch_summary(batch, include_changes=False)


async def list_batches(db, dataset_name: str | None = None, limit: int = 20) -> list[dict]:
    query = {"dataset": dataset_name} if dataset_name else {}
    cursor = db[UPLOADS].find(query).sort("uploaded_at", -1).limit(limit)
    return [batch_summary(b, include_changes=False) for b in await cursor.to_list(length=limit)]


async def get_batch_summary(db, upload_id: str) -> dict:
    return batch_summary(await _get_batch(db, upload_id))
