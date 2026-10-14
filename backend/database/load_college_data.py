"""
Loads the reviewed college data (IntelliVoice_College_Data_Review.xlsx) into MongoDB.

Run from the backend folder, with the venv active:

    python database/load_college_data.py "Z:\\intellivoice_data\\IntelliVoice_College_Data_Review.xlsx"
        -> DRY RUN: only prints what would be loaded. Nothing is written.

    python database/load_college_data.py "<xlsx>" --apply
        -> writes the records.

    python database/load_college_data.py --remove
        -> deletes everything this script ever added (and nothing else).

Safety rules
------------
* Every record it adds is tagged  _meta.seed_tag = "svit_college_data".
  Re-running --apply first removes the old tagged records of that
  collection, so running it twice never creates duplicates.
* A collection that already has records NOT added by this script
  (e.g. Bhuvana's fees, cutoffs, admissions, branches, recruiters) is
  skipped completely, unless you add --force.
* Rows whose status is "delete" or "CHECK", and empty values, are not loaded.
* The keys (.env) are read from your own .env file; this script never prints them.

The knowledge engine's cache refreshes within 60 s, so new data shows up
in IntelliVoice about a minute after --apply (no restart needed).
"""

import argparse
import os
import sys
from datetime import datetime, timezone

from dotenv import load_dotenv
from openpyxl import load_workbook
from pymongo import MongoClient

SEED_TAG = "svit_college_data"
SKIP_SHEETS = {"Read me", "Summary", "Questions for San", "New question types", "Profile updates"}
LOADABLE = {"from_file", "from_website", "from_san", "demo_sample", "system"}
REAL = {"from_file", "from_website", "from_san"}


def collection_name(sheet_title: str) -> str:
    return sheet_title.replace("(NEW)", "").strip()


def convert(value, typ):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if typ == "number":
        return int(float(str(value).replace(",", "")))
    if typ == "list":
        return [v.strip() for v in str(value).split(";") if v.strip()]
    return str(value).strip()


def read_workbook(path: str) -> dict:
    """{collection: [document, ...]} built from the review sheets."""
    wb = load_workbook(path, data_only=True)
    result = {}
    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS:
            continue
        docs = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            doc_id, field, value, typ, status, source = (list(row) + [None] * 6)[:6]
            status = (status or "").strip()
            if status not in LOADABLE:
                continue
            value = convert(value, (typ or "text").strip())
            if value is None:
                continue
            doc = docs.setdefault(doc_id, {"fields": {}, "demo": [], "sources": set()})
            doc["fields"][str(field).strip()] = value
            if status == "demo_sample":
                doc["demo"].append(str(field).strip())
            if status in REAL and source:
                doc["sources"].update(s.strip() for s in str(source).split(";") if s.strip())
        documents = []
        for doc_id, d in docs.items():
            if not any(k for k in d["fields"]):
                continue
            documents.append({
                **d["fields"],
                "_meta": {                      # never shown to the model (response_engine drops "_" keys)
                    "seed_tag": SEED_TAG,
                    "doc_id": doc_id,
                    "data_status": "contains_demo_sample" if d["demo"] else "verified",
                    "demo_fields": d["demo"],
                    "sources": sorted(d["sources"]),
                    "loaded_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                },
            })
        if documents:
            result[collection_name(ws.title)] = documents
    return result


def connect():
    load_dotenv()
    url = os.getenv("MONGODB_URL")
    if not url:
        sys.exit("MONGODB_URL not found in .env — run this from the backend folder.")
    client = MongoClient(url, serverSelectionTimeoutMS=15000)
    return client[os.getenv("DATABASE_NAME", "intellivoice")]


def main():
    parser = argparse.ArgumentParser(description="Load reviewed SVIT college data into MongoDB.")
    parser.add_argument("xlsx", nargs="?", help="path to IntelliVoice_College_Data_Review.xlsx")
    parser.add_argument("--apply", action="store_true", help="actually write (default is a dry run)")
    parser.add_argument("--remove", action="store_true", help="delete everything this script added")
    parser.add_argument("--force", action="store_true", help="also load into collections that already have other data")
    args = parser.parse_args()

    if args.remove:
        db = connect()
        total = 0
        for name in db.list_collection_names():
            n = db[name].delete_many({"_meta.seed_tag": SEED_TAG}).deleted_count
            if n:
                print(f"  removed {n:3d} from {name}")
                total += n
        print(f"Done. Removed {total} records added by this script.")
        return

    if not args.xlsx:
        parser.error("give the path to the review .xlsx (or use --remove)")

    data = read_workbook(args.xlsx)
    db = connect()
    print(("APPLYING" if args.apply else "DRY RUN (nothing will be written — add --apply)") + "\n")
    loaded = skipped = 0
    for name, docs in data.items():
        others = db[name].count_documents({"_meta.seed_tag": {"$ne": SEED_TAG}})
        demo = sum(1 for d in docs if d["_meta"]["demo_fields"])
        if others and not args.force:
            print(f"  SKIP  {name:24s} already has {others} records not added by this script")
            skipped += 1
            continue
        print(f"  {'LOAD' if args.apply else 'would load'}  {name:24s} {len(docs):2d} records"
              + (f"  ({demo} with demo values)" if demo else ""))
        if args.apply:
            db[name].delete_many({"_meta.seed_tag": SEED_TAG})
            db[name].insert_many(docs)
        loaded += len(docs)
    print(f"\n{'Loaded' if args.apply else 'Would load'} {loaded} records; skipped {skipped} collections.")
    if args.apply:
        print("IntelliVoice will use the new data within about a minute.")


if __name__ == "__main__":
    main()
